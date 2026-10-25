"""Tests for the offer-sheet import (api/catalogue_import.py).

The sample workbook is built here, shaped like Marginpar's Oct–Dec 2026 offer
sheet: terms above the table, a group-header row (Product / VBN / Grade / Color /
Standard Box / 1/8th box / an events block / Photos), a sub-header row, merged
group headers, hyperlinked photo cells, and a summary row that is not a single
variety.
"""

import os
from io import BytesIO

import frappe
import openpyxl
from frappe.tests.utils import FrappeTestCase

from upande_crm.api import catalogue_import as I

REAL_SHEET = "/home/austin/Downloads/Marginpar FOB Combined pricelist Oct-Dec 2026 final.xlsx"

ROWS = [
    # product, vbn, grade, color, std pack, bunch, price, tier, weight, eighth pack, eighth price, xmas, url
    ("Agapanthus 'Gletsjer'", 114199, "50 cm", "White", 320, "x10", 0.1856, 0.174, 11.28, None, None, 0.232,
     "https://www.marginpar.com/flowers/agapanthus-gletsjer"),
    ("Agapanthus 'Gletsjer'", 114199, "60 cm", "White", 300, "x10", 0.232, 0.2204, 11.28, 50, 0.2668, 0.29,
     "https://www.marginpar.com/flowers/agapanthus-gletsjer"),
    ("Aster (5 varieties)", None, "50 cm", "All", 800, "x25", 0.1392, 0.1276, 9.7, None, None, None,
     "https://www.marginpar.com/flowers"),
    ("Aster (Novi Belgii Grp) Dark Milka", 15557, "70 cm", "Purple", 400, "x25", 0.2088, 0.1972, 10.7, 50,
     0.2436, 0.3045, "https://www.marginpar.com/flowers/aster-dark-milka"),
]


def _tab(wb, title, factor):
    ws = wb.create_sheet(title)
    ws["B3"] = "Basic Lead Time:   5 Days From Day of Order"
    ws["B8"] = "Boxes we use: We use Marginpar branded boxes only"
    ws["B13"] = "Version date: 25 August 2026"
    ws["F13"] = "Fob Prices Sept-Dec 2026"
    ws["M13"] = "events prices"
    for col, val in {"B": "Product", "C": "VBN", "D": "Grade", "E": "Color", "F": "Standard Box",
                     "K": " 1/8th box", "M": "Standard Box", "N": " 1/8th box", "O": "Photos (Website)"}.items():
        ws[f"{col}14"] = val
    ws.merge_cells("F14:J14")
    ws.merge_cells("K14:L14")
    for col, val in {"F": "Packrate ", "G": "Bunch ", "H": " Price ", "I": "Price >4 boxes", "J": "Weight ",
                     "K": "Packrate", "L": "Price", "M": "Xmass/Saints Prices"}.items():
        ws[f"{col}15"] = val
    ws["J16"] = "100x33x20 (cm)"
    for i, r in enumerate(ROWS):
        n = 17 + i
        prod, vbn, grade, color, pack, bunch, price, tier, weight, epack, eprice, xmas, url = r
        for col, val in zip("BCDEFGHIJKLM", (prod, vbn, grade, color, pack, bunch,
                                              price * factor, tier * factor, weight, epack,
                                              eprice * factor if eprice else None, xmas)):
            ws[f"{col}{n}"] = val
        ws[f"O{n}"] = prod + " | Marginpar"
        ws[f"O{n}"].hyperlink = url


def sample_sheet():
    wb = openpyxl.Workbook()
    wb.active.title = "FOB Offersheet Oct-Dec 2026 fin"
    _tab(wb, "USD", 1)
    _tab(wb, "Euro", 0.86)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _line(lines, product, grade, box_type="Standard"):
    return next(l for l in lines if l["product"] == product and l["grade"] == grade and l["box_type"] == box_type)


class TestParse(FrappeTestCase):
    def test_parses_standard_and_eighth_box_lines(self):
        lines = I.parse_offer_sheet(sample_sheet(), ["USD", "EUR"])["lines"]
        std = _line(lines, "Agapanthus 'Gletsjer'", "60 cm")
        eighth = _line(lines, "Agapanthus 'Gletsjer'", "60 cm", "1/8 box")
        self.assertEqual((std["pack_qty"], std["bunch"], std["weight"], std["color"]), (300, "x10", 11.28, "White"))
        self.assertAlmostEqual(std["prices"]["USD"]["price"], 0.232, 3)
        self.assertAlmostEqual(std["prices"]["USD"]["tier"], 0.2204, 4)
        self.assertEqual(eighth["pack_qty"], 50)
        self.assertEqual(eighth["bunch"], "x10")
        self.assertAlmostEqual(eighth["prices"]["USD"]["price"], 0.2668, 4)
        self.assertNotIn("tier", eighth["prices"]["USD"])
        self.assertEqual(std["info_url"], "https://www.marginpar.com/flowers/agapanthus-gletsjer")
        self.assertEqual(std["vbn"], "114199")

    def test_no_eighth_line_without_eighth_price(self):
        lines = I.parse_offer_sheet(sample_sheet(), ["USD"])["lines"]
        self.assertFalse([l for l in lines if l["grade"] == "50 cm" and l["box_type"] == "1/8 box"])

    def test_events_block_is_ignored(self):
        lines = I.parse_offer_sheet(sample_sheet(), ["USD"])["lines"]
        std = _line(lines, "Aster (Novi Belgii Grp) Dark Milka", "70 cm")
        self.assertAlmostEqual(std["prices"]["USD"]["price"], 0.2088, 4)  # not the 0.3045 Xmas price

    def test_joins_usd_and_eur_tabs(self):
        lines = I.parse_offer_sheet(sample_sheet(), ["USD", "EUR"])["lines"]
        std = _line(lines, "Agapanthus 'Gletsjer'", "60 cm")
        self.assertAlmostEqual(std["prices"]["EUR"]["price"], 0.232 * 0.86, 4)
        self.assertEqual(len(lines), 6)  # 4 standard + 2 eighth, not doubled by the second tab

    def test_only_requested_currencies(self):
        lines = I.parse_offer_sheet(sample_sheet(), ["USD"])["lines"]
        self.assertTrue(all(set(l["prices"]) == {"USD"} for l in lines))

    def test_terms_come_from_above_the_table(self):
        terms = I.parse_offer_sheet(sample_sheet(), ["USD"])["terms"]
        self.assertIn("5 Days From Day of Order", terms)
        self.assertIn("branded boxes", terms)
        self.assertNotIn("Version date", terms)

    def test_no_currency_tab_raises(self):
        with self.assertRaises(frappe.ValidationError):
            I.parse_offer_sheet(sample_sheet(), ["GBP"])

    def test_real_sheet_smoke(self):
        if not os.path.exists(REAL_SHEET):
            self.skipTest("real offer sheet not on this machine")
        with open(REAL_SHEET, "rb") as fh:
            out = I.parse_offer_sheet(fh.read(), ["USD", "EUR"])
        self.assertGreater(len(out["lines"]), 100)
        self.assertTrue(all(l["prices"] for l in out["lines"]))


class TestMatch(FrappeTestCase):
    def tearDown(self):
        frappe.db.rollback()

    def _item(self, name):
        group = frappe.db.get_value("Item Group", {"is_group": 0}, "name")
        return frappe.get_doc({"doctype": "Item", "item_code": name, "item_name": name, "item_group": group,
                               "stock_uom": "Nos", "is_stock_item": 0}).insert(ignore_permissions=True).name

    def test_exact_and_normalised_names_match(self):
        milka = self._item("Aster (Novi Belgii Grp) Dark Milka TEST")
        glet = self._item("Agapanthus Gletsjer TEST")
        rows = [{"product": "Aster (Novi Belgii Grp) Dark Milka TEST", "vbn": None},
                {"product": "Agapanthus 'Gletsjer' test", "vbn": None}]
        out = I.match_items(rows)
        self.assertEqual((out[0]["item"], out[0]["match"]), (milka, "name"))
        self.assertEqual((out[1]["item"], out[1]["match"]), (glet, "normalised"))

    def test_summary_row_unmatched(self):
        out = I.match_items([{"product": "Aster (5 varieties) nowhere", "vbn": None}])
        self.assertIsNone(out[0]["item"])
        self.assertIsNone(out[0]["match"])
