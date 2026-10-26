"""Read a seasonal offer sheet (.xlsx) into catalogue lines.

Shaped on Marginpar's FOB offer sheet, but found by labels rather than fixed
cells, so a re-arranged sheet still reads:

* one tab per currency (`USD`, `Euro`/`EUR`); tabs that are not a currency the
  catalogue uses are ignored;
* a group-header row holding `Product`, `VBN`, `Grade`, `Color`, then block
  headers — `Standard Box`, `1/8th box`, a second `Standard Box` for event
  prices, `Photos (Website)`;
* a sub-header row under it (`Packrate`, `Bunch`, `Price`, `Price >4 boxes`,
  `Weight`).

The first Standard block and the first 1/8 block after it are read; anything
after them (event and Christmas prices) is ignored, as the spec decides. Each
block becomes its own line, so a row priced for both box types gives two.
Rows from every currency tab join on (product, grade, box type).

Matching to Items never creates one: unmatched rows go back to the user to map
by hand or skip.
"""

import json
import re

import frappe
from frappe import _
from frappe.utils import flt

from upande_crm.api.crm import _guard
from upande_crm.modules import requires_module

TAB_CURRENCY = {"USD": "USD", "EUR": "EUR", "EURO": "EUR", "GBP": "GBP", "KES": "KES"}
STANDARD, EIGHTH = "Standard", "1/8 box"


def _text(v):
    return re.sub(r"\s+", " ", str(v)).strip() if v not in (None, "") else ""


def _key(v):
    return _text(v).lower()


def _tab_currency(title, wanted):
    cur = TAB_CURRENCY.get(_text(title).upper()) or _text(title).upper()
    return cur if cur in wanted else None


def _header_row(ws):
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 60)):
        labels = {_key(c.value): c.column for c in row if c.value is not None}
        if {"product", "grade"} <= set(labels):
            return row[0].row, labels
    return None, None


def _blocks(ws, hrow, labels):
    """[(box_type, {field: column})] for the first Standard and 1/8 blocks."""
    heads = sorted((c.column, _key(c.value)) for c in ws[hrow] if c.value is not None)
    color_col = labels.get("color", labels["grade"])
    groups = [(col, txt) for col, txt in heads if col > color_col]
    out, seen = [], set()
    for i, (col, txt) in enumerate(groups):
        end = groups[i + 1][0] if i + 1 < len(groups) else ws.max_column + 1
        kind = STANDARD if "standard" in txt else EIGHTH if ("1/8" in txt or "eighth" in txt) else None
        if not kind:
            continue
        if kind in seen or (kind == STANDARD and EIGHTH in seen):
            break  # a second Standard block is the events price list
        seen.add(kind)
        cols = {}
        for c in range(col, end):
            sub = _key(ws.cell(row=hrow + 1, column=c).value)
            if sub.startswith("packrate"):
                cols.setdefault("pack_qty", c)
            elif sub.startswith("bunch"):
                cols.setdefault("bunch", c)
            elif sub.startswith("weight"):
                cols.setdefault("weight", c)
            elif sub.startswith("price") and ">" in sub:
                cols.setdefault("tier", c)
            elif sub == "price":
                cols.setdefault("price", c)
        out.append((kind, cols))
    return out


def _terms(ws, hrow):
    lines = []
    for row in ws.iter_rows(min_row=1, max_row=hrow - 1):
        for c in row:
            t = _text(c.value)
            if t and not t.lower().startswith("version date") and isinstance(c.value, str) and len(t) > 12:
                lines.append(t)
    return "".join(f"<p>{frappe.utils.escape_html(t)}</p>" for t in dict.fromkeys(lines))


def _url(cell):
    if cell is None:
        return None
    if cell.hyperlink and cell.hyperlink.target:
        return cell.hyperlink.target
    t = _text(cell.value)
    return t if t.startswith("http") else None


def parse_offer_sheet(content, currencies, spec_headers=None):
    """{"lines": [...], "terms": html} from the workbook bytes. Pure: no DB.

    `spec_headers` ({spec key: [sheet header, ...]}) reads a template's spec
    columns too, from the group-header row; each line then carries `specs`.
    """
    import openpyxl
    from io import BytesIO

    wanted = [c.upper() for c in currencies]
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    lines, order, terms = {}, [], ""
    tabs = [(ws, _tab_currency(ws.title, wanted)) for ws in wb.worksheets]
    tabs = [(ws, cur) for ws, cur in tabs if cur]
    if not tabs:
        frappe.throw(_("No tab in this workbook is named after {0}.").format(", ".join(wanted)))
    for ws, cur in tabs:
        hrow, labels = _header_row(ws)
        if not hrow:
            continue
        terms = terms or _terms(ws, hrow)
        blocks = _blocks(ws, hrow, labels)
        photo_col = next((c for k, c in labels.items() if k.startswith("photo")), None)
        spec_cols = {key: next((labels[_key(h)] for h in heads if _key(h) in labels), None)
                     for key, heads in (spec_headers or {}).items()}
        for r in range(hrow + 2, ws.max_row + 1):
            product = _text(ws.cell(row=r, column=labels["product"]).value)
            if not product:
                continue
            get = lambda col: ws.cell(row=r, column=col).value if col else None  # noqa: E731
            vbn = get(labels.get("vbn"))
            base = {"product": product, "grade": _text(get(labels["grade"])), "color": _text(get(labels.get("color"))),
                    "vbn": _text(int(vbn) if isinstance(vbn, float) and vbn.is_integer() else vbn) or None,
                    "info_url": _url(ws.cell(row=r, column=photo_col)) if photo_col else None}
            std_bunch = None
            for kind, cols in blocks:
                price = get(cols.get("price"))
                if kind == STANDARD:
                    std_bunch = _text(get(cols.get("bunch"))) or None
                if not isinstance(price, (int, float)) or price <= 0:
                    continue
                key = (product, base["grade"], kind)
                line = lines.get(key)
                if not line:
                    pack = get(cols.get("pack_qty"))
                    weight = get(cols.get("weight"))
                    line = {**base, "box_type": kind,
                            **({"specs": {k: _text(get(c)) for k, c in spec_cols.items() if c and _text(get(c))}}
                               if spec_headers is not None else {}),
                            "pack_qty": flt(pack) if isinstance(pack, (int, float)) else None,
                            "bunch": _text(get(cols.get("bunch"))) or std_bunch,
                            "weight": flt(weight) if isinstance(weight, (int, float)) else None,
                            "prices": {}}
                    lines[key] = line
                    order.append(key)
                entry = {"price": round(flt(price), 6)}
                tier = get(cols.get("tier"))
                if isinstance(tier, (int, float)) and tier > 0:
                    entry["tier"] = round(flt(tier), 6)
                line["prices"][cur] = entry
    return {"lines": [lines[k] for k in order], "terms": terms}


def _norm(name):
    return re.sub(r"\s+", " ", re.sub(r"[\"'‘’“”()]", " ", (name or "").lower())).strip()


def match_items(rows):
    """Add `item` and `match` ("vbn" | "name" | "normalised" | None) to each row."""
    meta = frappe.get_meta("Item")
    vbn_field = next((f for f in ("vbn", "vbn_code", "custom_vbn", "custom_vbn_code") if meta.has_field(f)), None)
    fields = ["name", "item_name"] + ([vbn_field] if vbn_field else [])
    items = frappe.get_all("Item", filters={"disabled": 0}, fields=fields, limit=0)
    by_vbn = {_text(i.get(vbn_field)): i.name for i in items if vbn_field and i.get(vbn_field)}
    by_name, by_norm = {}, {}
    for i in items:
        for n in {i.name, i.item_name}:
            if n:
                by_name.setdefault(_key(n), i.name)
                by_norm.setdefault(_norm(n), i.name)
    for row in rows:
        vbn = _text(row.get("vbn"))
        if vbn and vbn in by_vbn:
            row["item"], row["match"] = by_vbn[vbn], "vbn"
        elif _key(row["product"]) in by_name:
            row["item"], row["match"] = by_name[_key(row["product"])], "name"
        elif _norm(row["product"]) in by_norm:
            row["item"], row["match"] = by_norm[_norm(row["product"])], "normalised"
        else:
            row["item"], row["match"] = None, None
    return rows


@frappe.whitelist(methods=["POST"])
@requires_module("cat")
def crm_catalogue_import(file_url, currencies, catalogue=None, template=None):
    """Preview an uploaded offer sheet as catalogue lines; nothing is saved.

    With a template (or a catalogue, whose template is used) the template's
    spec columns are read from the sheet headers named in `import_headers`:
    Column specs stay on the lines, Badge/Subtitle specs come back per item in
    `products`.
    """
    _guard()
    if isinstance(currencies, str):
        currencies = json.loads(currencies) if currencies.startswith("[") else currencies.split(",")
    currencies = [c.strip().upper() for c in currencies if c.strip()]
    f = frappe.get_doc("File", {"file_url": file_url})
    f.check_permission("read")
    template = template or (frappe.db.get_value("CRM Catalogue", catalogue, "template") if catalogue else None)
    specs = frappe.get_doc("CRM Catalogue Template", template).spec_columns if template else []
    headers = {s.key: [h.strip() for h in (s.import_headers or s.label).split(",") if h.strip()] for s in specs}
    out = parse_offer_sheet(f.get_content(), currencies, spec_headers=headers if template else None)
    rows = match_items(out["lines"])
    products = {}
    if template:
        product_keys = {s.key for s in specs if s.display != "Column"}
        for r in rows:
            found = r.get("specs") or {}
            r["specs"] = {k: v for k, v in found.items() if k not in product_keys}
            if r["item"]:
                products.setdefault(r["item"], {}).update({k: v for k, v in found.items() if k in product_keys})
    return {"lines": [r for r in rows if r["item"]], "unmatched": [r for r in rows if not r["item"]],
            "terms": out["terms"],
            "products": [{"item": k, "specs": v} for k, v in products.items() if v]}
